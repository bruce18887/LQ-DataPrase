# test/backend/test_weight_xlsx.py
"""apps/weight/xlsx.py 的**工作簿结构**：总览表、超链接、每测试项一 sheet 与明细数据。

2026-10-20 用户点定：导出形态对齐批量参数分布图 ——
`总览` 列测试项（超链接），每个可见测试项一张 sheet，图在 A1、数据在下方、A23 返回总览。
图本身的圆点/标签/坐标轴在 `test_weight_xlsx_chart.py`；本文件只断结构、数据与表名。

夹具形状：56 声明 = 14 组 × 4、长度 {4,8,16,32,64}、小数位全 5。
「按原文小数位写 number_format」「均值 6 位」仍用合成/手工文档断（夹具写死 5 位也绿）。
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings.development')
import django  # noqa: E402
django.setup()
from django.test import TestCase  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from apps.weight.parser import (  # noqa: E402
    WeightDoc, WeightGroup, WeightSeries, compute_mean, parse_weight_text,
)
from apps.weight.xlsx import (  # noqa: E402
    DEFAULT_PALETTE, DETAIL_COL_NAME, DETAIL_COL_SOURCE_INDEX, DETAIL_COL_VALUE0,
    DETAIL_FIRST_DATA_ROW, DETAIL_HEADER_ROW, DETAIL_RETURN_ROW, SUMMARY_HEADERS,
    SUMMARY_SHEET, build_weight_workbook, weight_detail_sheet_names,
)

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIXTURE = os.path.join(_ROOT, 'test', 'fixtures', 'weight', 'trim_weight_ci.csv')
GROUPS = 14
PER_GROUP = 4
ALL_SERIES = 56
_MEMBERS = (1, 2, 3, 4)


def _doc():
    with io.open(FIXTURE, encoding='utf-8', newline='') as f:
        return parse_weight_text(f.read())


def _wb(doc=None, **kw):
    kw.setdefault('palette', DEFAULT_PALETTE)
    return load_workbook(io.BytesIO(build_weight_workbook(doc or _doc(), **kw)))


def _summary(wb):
    return wb[SUMMARY_SHEET]


def _group(index, name, lens=(4,), decimals=5):
    """合成组：成员长度按 lens，值取 0..n-1（便于逐格核对）。"""
    series = [
        WeightSeries(f'{name}_{k + 1}', k + 1, n, decimals,
                     tuple(float(i) for i in range(n)))
        for k, n in enumerate(lens)
    ]
    return WeightGroup(index=index, name=name, first_row=1,
                       series=tuple(series), mean=compute_mean(series, f'{name}_Avg'))


def _flat(doc):
    return [s for g in doc.groups for s in g.series]


class SummaryTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.doc = _doc()
        cls.wb = _wb(cls.doc)
        cls.ws = _summary(cls.wb)

    def test_sheet_list_is_summary_then_visible_groups(self):
        self.assertEqual(self.wb.sheetnames,
                         [SUMMARY_SHEET] + [f'WT{i + 1:04d}' for i in range(GROUPS)])

    def test_header_row(self):
        self.assertEqual([self.ws.cell(1, c).value for c in range(1, len(SUMMARY_HEADERS) + 1)],
                         list(SUMMARY_HEADERS))

    def test_one_row_per_group_in_first_appearance_order(self):
        self.assertEqual([self.ws.cell(r, 1).value for r in range(2, 16)],
                         list(range(1, GROUPS + 1)))
        self.assertEqual([self.ws.cell(r, 2).value for r in range(2, 16)],
                         [g.name for g in self.doc.groups])

    def test_test_item_cell_is_a_hyperlink_to_its_sheet(self):
        for row, g in enumerate(self.doc.groups, start=2):
            cell = self.ws.cell(row, 2)
            self.assertEqual(cell.hyperlink.target, f"#'{g.name}'!A1", g.name)
            self.assertTrue(str(cell.font.color.rgb).endswith('0563C1'), g.name)

    def test_row_stats(self):
        # 对照组：先确认夹具的组长与点数真的不同，避免「全 4/64」恒真
        self.assertEqual([max(len(s.values) for s in g.series) for g in self.doc.groups],
                         [16, 8, 16, 8, 64, 16, 32, 16, 4, 4, 8, 32, 16, 32])
        first = [self.ws.cell(2, c).value for c in (3, 4, 5)]
        self.assertEqual(first, [4, 64, 15])
        fifth = [self.ws.cell(6, c).value for c in (3, 4, 5)]
        self.assertEqual(fifth, [4, 256, 63])

    def test_min_max_are_over_all_raw_values(self):
        raw = [v for s in self.doc.groups[0].series for v in s.values]
        self.assertEqual(self.ws.cell(2, 6).value, min(raw))
        self.assertEqual(self.ws.cell(2, 7).value, max(raw))
        self.assertEqual((self.ws.cell(2, 6).value, self.ws.cell(2, 7).value),
                         (-0.02774, 0.03098))

    def test_column_widths(self):
        widths = [self.ws.column_dimensions[get_column_letter(c)].width for c in range(1, 8)]
        self.assertEqual(widths, [8, 40, 10, 12, 12, 16, 16])


class VisibleFilterTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.doc = _doc()

    def test_only_visible_groups_get_sheets_and_rows(self):
        wb = _wb(self.doc, visible_group_indexes=[1, 3])
        self.assertEqual(wb.sheetnames, [SUMMARY_SHEET, 'WT0002', 'WT0004'])
        ws = _summary(wb)
        self.assertEqual([ws.cell(r, 2).value for r in range(2, 4)], ['WT0002', 'WT0004'])
        self.assertIsNone(ws.cell(4, 1).value)

    def test_hidden_group_data_is_not_exported(self):
        wb = _wb(self.doc, visible_group_indexes=[1])
        self.assertNotIn('WT0001', wb.sheetnames)
        self.assertIn('WT0002', wb.sheetnames)

    def test_empty_visibility_keeps_only_summary_header(self):
        wb = _wb(self.doc, visible_group_indexes=[])
        self.assertEqual(wb.sheetnames, [SUMMARY_SHEET])
        ws = _summary(wb)
        self.assertEqual([ws.cell(2, c).value for c in range(1, 8)], [None] * 7)
        self.assertEqual(len(wb[SUMMARY_SHEET]._charts), 0)

    def test_junk_indexes_degrade_to_no_details(self):
        wb = _wb(self.doc, visible_group_indexes=[-1, 99, 'x', None, {}])
        self.assertEqual(wb.sheetnames, [SUMMARY_SHEET])


class DetailSheetTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.doc = _doc()
        cls.wb = _wb(cls.doc)

    def _detail(self, name='WT0001'):
        return self.wb[name]

    def test_header_row_and_code_columns(self):
        ws = self._detail()
        self.assertEqual(ws.cell(DETAIL_HEADER_ROW, DETAIL_COL_SOURCE_INDEX).value, '原序号')
        self.assertEqual(ws.cell(DETAIL_HEADER_ROW, DETAIL_COL_NAME).value, '系列名')
        self.assertEqual([ws.cell(DETAIL_HEADER_ROW, DETAIL_COL_VALUE0 + c).value for c in range(16)],
                         list(range(16)))
        self.assertEqual(ws.cell(DETAIL_HEADER_ROW, DETAIL_COL_VALUE0 + 15).value, 15)
        self.assertIsNone(ws.cell(DETAIL_HEADER_ROW, DETAIL_COL_VALUE0 + 16).value)

    def test_four_member_rows_and_source_indexes(self):
        ws = self._detail()
        names = [ws.cell(DETAIL_FIRST_DATA_ROW + i, DETAIL_COL_NAME).value for i in range(4)]
        self.assertEqual(names, [f'WT0001_{k}' for k in _MEMBERS])
        self.assertEqual([ws.cell(DETAIL_FIRST_DATA_ROW + i, DETAIL_COL_SOURCE_INDEX).value
                          for i in range(4)], [1, 15, 29, 43])

    def test_every_written_value_equals_the_parsed_one(self):
        """14 张分表逐条比：1088 个点一个不多一个不少。"""
        row = DETAIL_FIRST_DATA_ROW
        for g in self.doc.groups:
            ws = self.wb[g.name]
            row = DETAIL_FIRST_DATA_ROW
            for s in g.series:
                got = [ws.cell(row, DETAIL_COL_VALUE0 + c).value for c in range(len(s.values))]
                self.assertEqual(got, list(s.values), s.name)
                self.assertIsNone(ws.cell(row, DETAIL_COL_VALUE0 + len(s.values)).value, s.name)
                row += 1

    def test_number_format_follows_source_decimals(self):
        doc = parse_weight_text('double A[2] = {1.0, 2.5};\n'
                                'double B[2] = {3.25, 4.75};\n'
                                'double C[2] = {5, 6};\n')
        wb = _wb(doc)
        got = [wb['A'].cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_VALUE0).number_format,
               wb['B'].cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_VALUE0).number_format,
               wb['C'].cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_VALUE0).number_format]
        self.assertEqual(got, ['0.0', '0.00', '0'])

    def test_mean_row_follows_members_when_shown(self):
        wb = _wb(self.doc, show_mean=True)
        ws = wb['WT0001']
        row = DETAIL_FIRST_DATA_ROW + PER_GROUP
        self.assertEqual(ws.cell(row, DETAIL_COL_NAME).value, 'WT0001_Avg')
        self.assertIsNone(ws.cell(row, DETAIL_COL_SOURCE_INDEX).value)
        self.assertTrue(ws.cell(row, DETAIL_COL_NAME).font.bold)
        self.assertEqual(ws.cell(row, DETAIL_COL_VALUE0 + 3).number_format, '0.000000')
        self.assertEqual(ws.cell(row, DETAIL_COL_VALUE0 + 3).value, -0.0119525)

    def test_mean_rows_absent_by_default(self):
        ws = self._detail()
        self.assertNotIn('WT0001_Avg',
                         [ws.cell(DETAIL_FIRST_DATA_ROW + i, DETAIL_COL_NAME).value for i in range(4)])

    def test_return_link_at_a23(self):
        ws = self._detail()
        cell = ws.cell(DETAIL_RETURN_ROW, DETAIL_COL_SOURCE_INDEX)
        self.assertEqual(cell.value, '← 返回总览')
        self.assertEqual(cell.hyperlink.target, f"#'{SUMMARY_SHEET}'!A1")
        self.assertTrue(str(cell.font.color.rgb).endswith('0563C1'))

    def test_group_with_64_codes_uses_its_own_width(self):
        ws = self.wb['WT0005']
        last = DETAIL_COL_VALUE0 + 63
        self.assertEqual(ws.cell(DETAIL_HEADER_ROW, last).value, 63)
        self.assertIsNone(ws.cell(DETAIL_HEADER_ROW, last + 1).value)
        self.assertEqual(ws.column_dimensions[get_column_letter(last)].width, 12)

    def test_name_cell_colours_and_alternating_fill(self):
        even = self.wb['WT0001']   # index 0：不铺底
        odd = self.wb['WT0002']    # index 1：极浅底
        self.assertIsNone(even.cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_NAME).fill.patternType)
        self.assertEqual(odd.cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_NAME).fill.patternType, 'solid')
        self.assertTrue(str(odd.cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_NAME).font.color.rgb)
                        .endswith(DEFAULT_PALETTE[1]))
        self.assertTrue(str(even.cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_NAME).font.color.rgb)
                        .endswith(DEFAULT_PALETTE[0]))

    def test_each_detail_sheet_has_exactly_one_chart_and_summary_has_none(self):
        self.assertEqual(self.wb[SUMMARY_SHEET]._charts, [])
        for g in self.doc.groups:
            self.assertEqual(len(self.wb[g.name]._charts), 1, g.name)


class AbsentAndEmptyTests(TestCase):
    def test_absent_point_stays_blank(self):
        a = WeightSeries('A_1', 1, 3, 1, (2.5, None, 3.5))
        b = WeightSeries('A_2', 2, 1, 0, (4.5,))
        g = WeightGroup(index=0, name='A', first_row=1, series=(a, b),
                        mean=compute_mean((a, b), 'A_Avg'))
        doc = WeightDoc(filename_hint='syn', groups=(g,), code_max=2, series_count=2, group_count=1)
        ws = _wb(doc, show_mean=True)['A']
        self.assertEqual(ws.cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_VALUE0).value, 2.5)
        self.assertIsNone(ws.cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_VALUE0 + 1).value)
        self.assertEqual(ws.cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_VALUE0 + 2).value, 3.5)

    def test_zero_length_group_still_gets_a_sheet_but_no_chart(self):
        doc = parse_weight_text('double E[0] = {};\ndouble F[2] = {1.5, 2.5};\n')
        wb = _wb(doc)
        self.assertEqual(wb.sheetnames, [SUMMARY_SHEET, 'E', 'F'])
        self.assertEqual(len(wb['E']._charts), 0)          # 无点不建图
        self.assertEqual(wb['E'].cell(DETAIL_FIRST_DATA_ROW, DETAIL_COL_NAME).value, 'E')
        self.assertEqual(len(wb['F']._charts), 1)
        ws = _summary(wb)
        self.assertEqual([ws.cell(2, c).value for c in (3, 4, 5, 6, 7)],
                         [1, 0, 0, 'N/A', 'N/A'])


class SheetNameTests(TestCase):
    def test_illegal_chars_become_underscore(self):
        names = weight_detail_sheet_names([_group(0, "a[]:*?/\\b'")])
        self.assertEqual(names[0], 'a' + '_' * 7 + 'b' + '_')

    def test_long_names_are_truncated_to_31(self):
        names = weight_detail_sheet_names([_group(0, 'x' * 99)])
        self.assertEqual(len(names[0]), 31)

    def test_duplicate_names_get_numeric_suffixes(self):
        groups = [_group(0, 'A'), _group(1, 'A'), _group(2, 'A')]
        names = weight_detail_sheet_names(groups)
        self.assertEqual([names[i] for i in range(3)], ['A', 'A_2', 'A_3'])

    def test_collisions_are_case_insensitive(self):
        names = weight_detail_sheet_names([_group(0, 'Abc'), _group(1, 'abc')])
        self.assertEqual([names[0], names[1]], ['Abc', 'abc_2'])

    def test_summary_and_history_are_reserved(self):
        names = weight_detail_sheet_names([
            _group(0, SUMMARY_SHEET), _group(1, 'History'), _group(2, 'history')])
        self.assertEqual(names[0], f'{SUMMARY_SHEET}_2')
        self.assertEqual(names[1], 'History_2')
        self.assertEqual(names[2], 'history_3')

    def test_empty_name_falls_back(self):
        self.assertEqual(weight_detail_sheet_names([_group(0, '   ')])[0], '测试项')

    def test_truncated_collision_keeps_the_suffix(self):
        base = 'y' * 31
        names = weight_detail_sheet_names([_group(0, base), _group(1, base + 'z')])
        self.assertEqual(len(names[0]), 31)
        self.assertEqual(names[1], base[:29] + '_2')
        self.assertLessEqual(len(names[1]), 31)
