"""Buyoff Excel 版式缺陷回归测试（缺陷 #4 / #5 / #6 / #7）。

直接驱动 ``apps.buyoff.excelize_layout.build_buyoff_form``（真实 excelize 句柄），
把生成的单元格值、填充色、以及**条件格式规则**读回来断言：

- 判定语义由 Result 区的**条件格式**承载（`>fail` 红 / `>warn` 黄 / 其余浅绿基色）。
  excelize 没有 CF 读取接口，所以用例会把工作簿存下来用 openpyxl 读回规则，
  再按 `stop_if_true` 顺序**模拟 Excel 求值**，断言具体数值命中哪一档 —— 比只断言
  "规则存在"更能守住业务口径。
- 「无法判定」（N/A）是静态浅灰底，不参与条件格式。

- #4 ``qa_range == 0`` 时旧实现把分母换成「1 个工程单位」，Result 列输出无意义
  大数并被判红 → 应写 ``'N/A'`` + 灰色样式，且不参与红/黄/绿判定；
- #5 ``all_stats['FT'][param]`` 无守卫 → KeyError/500（视图会跳过 stats 为空的
  参数，FT 缺该参数是真实可能）；
- #6 限值解析失败被渲染成**红色 FAIL**（``diff = 0`` → ``diff <= 0``）——
  「无法判定」显示成「不合格」，语义相反 → 应写 ``'N/A'`` + 灰色；
- #7 百分比 2 位小数把 0.002% 归零 → 项目口径 6 位小数。

runner: ``manage.py test test.backend.test_buyoff_layout``
"""

import io
import os
import tempfile

import excelize
from django.test import SimpleTestCase
from openpyxl import load_workbook

from apps.buyoff import excelize_layout as layout
from apps.export import excel_theme
from test.backend.excel_cf import cf_rules_for_cell, cf_verdict

SHEET = 'Buyoff data'

# Result 区判定改由条件格式承载。判定格有静态浅绿基色，|pct| 过 warn/fail 时由 CF 覆盖。
NA_FILL = excel_theme.FILL_NA
BASE_FILL = excel_theme.FILL_PASS

PARAM = 'V_R'


def _stats(mean, lower=0.5, upper=1.5):
    """一份最小可用的 buyoff stats（模拟 compute_buyoff_stats 的输出）。"""
    return {
        'lower_limit': lower, 'upper_limit': upper,
        'min': mean - 0.1, 'max': mean + 0.1, 'range': 0.2,
        'mean': mean, 'std': 0.05,
        'mean_minus_6std': mean - 0.3, 'mean_minus_3std': mean - 0.15,
        'mean_plus_3std': mean + 0.15, 'mean_plus_6std': mean + 0.3,
        'ca': 0.0, 'cp': 3.33, 'cpk': 3.0,
    }


def _metadata(min_raw='0.5', max_raw='1.5', present=True):
    if not present:
        return {'format': 'CTA8290D', 'units': {}, 'mins': {}, 'maxs': {}}
    return {
        'format': 'CTA8290D',
        'units': {PARAM: 'mV'},
        'mins': {PARAM: min_raw},
        'maxs': {PARAM: max_raw},
    }


def _ds(min_raw='0.5', max_raw='1.5', file_id=1, present=True):
    return {'df': None, 'metadata': _metadata(min_raw, max_raw, present),
            'numeric_cols': [PARAM], 'file_id': file_id}


def _datasets(ft=('0.5', '1.5'), qa1=('0.5', '1.5'), qa2=('0.5', '1.5'),
              present=True):
    return {
        'FT.csv': _ds(*ft, 1, present),
        'QA1.csv': _ds(*qa1, 2, present),
        'QA2.csv': _ds(*qa2, 3, present),
    }


ROLE_MAPPING = {'FT': 'FT.csv', 'QA1': 'QA1.csv', 'QA2': 'QA2.csv'}


def _safe_close(f):
    try:
        f.close()
    except Exception:  # noqa: BLE001 — 已 close 过就无所谓
        pass


def _load_workbook(f):
    """存到临时文件再读回。用 save_as（不是 save_excelize）以免关闭 f，后续还能继续读单元格。"""
    tmp = tempfile.mktemp(suffix='.xlsx')
    try:
        f.save_as(tmp)
        with open(tmp, 'rb') as fh:
            data = fh.read()
    finally:
        try:
            os.unlink(tmp)
        except OSError:
            pass
    return load_workbook(io.BytesIO(data))


class _LayoutBase(SimpleTestCase):
    def _build(self, all_stats, datasets=None, role_mapping=None,
               ordered_roles=None, common_items=(PARAM,)):
        f = excelize.new_file()
        self.addCleanup(_safe_close, f)
        layout.build_buyoff_form(
            f,
            role_mapping if role_mapping is not None else dict(ROLE_MAPPING),
            list(common_items),
            all_stats,
            datasets if datasets is not None else _datasets(),
            ordered_roles if ordered_roles is not None else ['FT', 'QA1', 'QA2'],
        )
        return f

    def _label_rows(self, f):
        rows = {}
        for r in range(1, 120):
            val = f.get_cell_value(SHEET, f'B{r}')
            if val:
                rows[str(val)] = r
        return rows

    def _cell(self, f, label, col='C'):
        row = self._label_rows(f)[label]
        return f'{col}{row}'

    def _value(self, f, label, col='C'):
        # Result 区四行现在是**公式**：公式格 get_cell_value 读回空串，必须用
        # calc_cell_value。它返回**格式化后**的显示值，正好与下面这些
        # 'N/A' / 'x.xxxxxx%' 断言一致。
        return f.calc_cell_value(SHEET, self._cell(f, label, col))

    def _formula(self, f, label, col='C'):
        """取该格的公式串（不含前导 '='）；非公式格返回 ''。"""
        return f.get_cell_formula(SHEET, self._cell(f, label, col))

    def _fill(self, f, label, col='C'):
        cell = self._cell(f, label, col)
        style = f.get_style(f.get_cell_style(SHEET, cell))
        fill = getattr(style, 'fill', None)
        colors = getattr(fill, 'color', None) if fill is not None else None
        return str(colors[0]).upper() if colors else ''

    def _rules_for_cell(self, f, label, col='C'):
        """取覆盖该格的 CF 规则（按优先级排序），供 cf_verdict 模拟求值。"""
        ws = _load_workbook(f)[SHEET]
        return cf_rules_for_cell(ws, self._cell(f, label, col))


QA1_ROW = '(QA1 Mean - FT Mean)/ (QA Upper Limit - QA Lower Limit)'
QA2_ROW = '(QA2 Mean - FT Mean)/ (QA Upper Limit - QA Lower Limit)'
FT_LL_ROW = 'FT Lower Limit - QA Lower Limit'
QA_UL_ROW = 'QA Upper Limit - FT Upper Limit'


class ZeroQaRangeTests(_LayoutBase):
    """缺陷 #4：qa_range == 0 → 'N/A' + 灰色，不参与红/黄/绿判定。"""

    def test_zero_range_writes_na(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            # QA1 上下限相同 → 公差为 0，百分比无定义
            'QA1': {PARAM: _stats(1.5, lower=2.0, upper=2.0)},
            'QA2': {PARAM: _stats(1.5, lower=2.0, upper=2.0)},
        }
        f = self._build(all_stats)
        for label in (QA1_ROW, QA2_ROW):
            self.assertEqual(self._value(f, label), 'N/A', f'{label} 公差为 0 应写 N/A')
            self.assertEqual(self._fill(f, label), NA_FILL,
                             f'{label} 无法判定应是静态浅灰底')
            self.assertNotEqual(self._fill(f, label), BASE_FILL,
                                f'{label} 无法判定不得用判定基色（浅绿）')

    def test_missing_limit_writes_na(self):
        """限值缺失（None）→ 公差未知 → 同样 'N/A' + 灰色。"""
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.5, lower=None, upper=None)},
            'QA2': {PARAM: _stats(1.5, lower=None, upper=1.5)},
        }
        f = self._build(all_stats)
        for label in (QA1_ROW, QA2_ROW):
            self.assertEqual(self._value(f, label), 'N/A')
            self.assertEqual(self._fill(f, label), NA_FILL)

    def test_valid_range_still_judged(self):
        """正向对照：公差可用时仍按阈值给绿/黄/红。"""
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.01)},          # 1% → 绿
            'QA2': {PARAM: _stats(1.15)},          # 15% → 红
        }
        f = self._build(all_stats)
        # 1% 未过 5% 门槛 → 条件格式不介入，留在基色（浅绿）
        self.assertIsNone(cf_verdict(self._rules_for_cell(f, QA1_ROW), 0.01))
        self.assertEqual(self._fill(f, QA1_ROW), BASE_FILL)
        # 15% > 10% → 判红
        self.assertEqual(cf_verdict(self._rules_for_cell(f, QA2_ROW), 0.15),
                         excel_theme.FILL_FAIL)


class MissingFtStatTests(_LayoutBase):
    """缺陷 #5：FT 缺该 param 的 stats 时不得 KeyError。"""

    def test_ft_missing_param_does_not_raise(self):
        all_stats = {
            'FT': {},                       # 视图会跳过 stats 为空的参数
            'QA1': {PARAM: _stats(1.01)},
            'QA2': {PARAM: _stats(1.02)},
        }
        f = self._build(all_stats)   # 修复前：KeyError → 500
        for label in (QA1_ROW, QA2_ROW):
            self.assertEqual(self._value(f, label), 'N/A')
            self.assertEqual(self._fill(f, label), NA_FILL)

    def test_ft_role_absent_does_not_raise(self):
        """只有 FT+QA2（无 QA1 角色）时两行 Result 都不能崩。"""
        all_stats = {'FT': {PARAM: _stats(1.0)}, 'QA2': {PARAM: _stats(1.01)}}
        f = self._build(
            all_stats,
            role_mapping={'FT': 'FT.csv', 'QA2': 'QA2.csv'},
            ordered_roles=['FT', 'QA2'],
        )
        self.assertEqual(self._value(f, QA1_ROW), 'N/A')
        self.assertEqual(self._value(f, QA2_ROW), '1.000000%')


class UnparsableLimitRenderTests(_LayoutBase):
    """缺陷 #6：限值解析失败必须 'N/A' + 灰色，绝不能是红色 FAIL。"""

    def test_text_limits_render_na_not_red(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.0)},
            'QA2': {PARAM: _stats(1.0)},
        }
        f = self._build(all_stats, datasets=_datasets(ft=('Min', 'Max'),
                                                      qa1=('Min', 'Max'),
                                                      qa2=('Min', 'Max')))
        for label in (FT_LL_ROW, QA_UL_ROW):
            self.assertEqual(self._value(f, label), 'N/A',
                             f'{label} 限值不可解析应写 N/A')
            self.assertEqual(self._fill(f, label), NA_FILL,
                             f'{label} 限值不可解析应带 N/A 灰底')
            self.assertNotEqual(self._fill(f, label), excel_theme.FILL_FAIL,
                                f'{label}「无法判定」不得渲染成红色 FAIL')

    def test_missing_limits_render_na_not_red(self):
        """限值缺失（metadata 里根本没有该 param）同样不得判红。"""
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.0)},
            'QA2': {PARAM: _stats(1.0)},
        }
        f = self._build(all_stats, datasets=_datasets(present=False))
        for label in (FT_LL_ROW, QA_UL_ROW):
            self.assertEqual(self._value(f, label), 'N/A')
            self.assertEqual(self._fill(f, label), NA_FILL)

    def test_numeric_limits_still_judged_red_when_tighter(self):
        """正向对照：QA 下限高于 FT 下限（diff<=0）仍判红。"""
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.0)},
            'QA2': {PARAM: _stats(1.0)},
        }
        # FT 下限 0.5、QA1 下限 0.8 → diff = -0.3 → 红
        f = self._build(all_stats, datasets=_datasets(ft=('0.5', '1.5'),
                                                      qa1=('0.8', '1.5')))
        self.assertAlmostEqual(float(self._value(f, FT_LL_ROW)), -0.3, places=6)
        self.assertEqual(cf_verdict(self._rules_for_cell(f, FT_LL_ROW), -0.3),
                         excel_theme.FILL_FAIL)


class PercentagePrecisionTests(_LayoutBase):
    """缺陷 #7：百分比统一 6 位小数（1/50000 = 0.002% 不得归零）。"""

    def test_tiny_percentage_keeps_six_decimals(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            # mean_diff = 0.00002，qa_range = 1.0 → 0.002%
            'QA1': {PARAM: _stats(1.00002)},
            'QA2': {PARAM: _stats(1.00002)},
        }
        f = self._build(all_stats)
        self.assertEqual(self._value(f, QA1_ROW), '0.002000%')
        self.assertEqual(self._value(f, QA2_ROW), '0.002000%')
        # 微小偏移仍是绿（未超 5% 阈值）
        self.assertIsNone(cf_verdict(self._rules_for_cell(f, QA1_ROW), 0.00002))
        self.assertEqual(self._fill(f, QA1_ROW), BASE_FILL)

    def test_normal_percentage_has_six_decimals(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.06)},   # 6% → 黄
            'QA2': {PARAM: _stats(1.06)},
        }
        f = self._build(all_stats)
        self.assertEqual(self._value(f, QA1_ROW), '6.000000%')
        self.assertEqual(cf_verdict(self._rules_for_cell(f, QA1_ROW), 0.06),
                         excel_theme.FILL_WARN)


class NaStyleConstantTests(_LayoutBase):
    """「无法判定」必须与判定档视觉可区分：前者是静态浅灰底，后者走浅绿基色 + 条件格式。"""

    def test_na_and_verdict_are_distinguishable(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.5, lower=2.0, upper=2.0)},  # 公差 0 → 无法判定
            'QA2': {PARAM: _stats(1.01)},                        # 1% → 通过
        }
        f = self._build(all_stats)
        # 无法判定：静态浅灰底，且该行虽有条件格式，数值档都命中不了
        self.assertEqual(self._fill(f, QA1_ROW), NA_FILL)
        self.assertIsNone(cf_verdict(self._rules_for_cell(f, QA1_ROW), 0.0))
        # 可判定：基色浅绿 + 条件格式覆盖到该行
        self.assertEqual(self._fill(f, QA2_ROW), BASE_FILL)
        self.assertTrue(self._rules_for_cell(f, QA2_ROW))

    def test_stat_rows_render_na_for_none_limits(self):
        """Lower/Upper Limit 行：限值不可解析 → 'N/A'（不是空白/0）。"""
        all_stats = {
            'FT': {PARAM: dict(_stats(1.0), lower_limit=None, upper_limit=None,
                               ca='N/A', cp='N/A', cpk='N/A')},
            'QA1': {PARAM: _stats(1.0)},
            'QA2': {PARAM: _stats(1.0)},
        }
        f = self._build(all_stats, datasets=_datasets(ft=('Min', 'Max'),
                                                      qa1=('Min', 'Max'),
                                                      qa2=('Min', 'Max')))
        rows = self._label_rows(f)
        self.assertGreater(len(rows), 10)
        # FT 段的 Cpk 行（第一次出现）应为 N/A
        cpk_rows = [r for r in range(1, 120)
                    if str(f.get_cell_value(SHEET, f'B{r}')) == 'Cpk']
        self.assertTrue(cpk_rows)
        self.assertEqual(f.get_cell_value(SHEET, f'C{cpk_rows[0]}'), 'N/A')


class ResultFormulaTests(_LayoutBase):
    """Result 四行改为 Excel 公式：可审计、随限值/均值重算。

    公式只引用**本表**已写好的限值行与均值行（同表引用，不依赖外部原始数据），
    且仅在可判定时写入；不可判定仍是静态 'N/A'（不落公式）。
    """

    def test_result_rows_are_formulas(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.1)},
            'QA2': {PARAM: _stats(1.05)},
        }
        f = self._build(all_stats)
        for label in (FT_LL_ROW, QA_UL_ROW, QA1_ROW, QA2_ROW):
            self.assertTrue(self._formula(f, label).startswith('ROUND('),
                            f'{label} 应为公式，实际 {self._formula(f, label)!r}')

    def test_result_formula_values_match_hand_calc(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.1)},
            'QA2': {PARAM: _stats(1.05)},
        }
        f = self._build(all_stats, datasets=_datasets(ft=('0.5', '1.5'),
                                                      qa1=('0.8', '1.6'),
                                                      qa2=('0.5', '1.5')))
        self.assertAlmostEqual(float(self._value(f, FT_LL_ROW)), 0.5 - 0.8, places=6)  # -0.3
        self.assertAlmostEqual(float(self._value(f, QA_UL_ROW)), 1.6 - 1.5, places=6)  # 0.1
        self.assertEqual(self._value(f, QA1_ROW), '12.500000%')  # (1.1-1.0)/(1.6-0.8)
        self.assertEqual(self._value(f, QA2_ROW), '5.000000%')   # (1.05-1.0)/(1.5-0.5)

    def test_na_cells_carry_no_formula(self):
        """不可判定（公差 0）时写静态 'N/A'，不落公式。"""
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.5, lower=2.0, upper=2.0)},
            'QA2': {PARAM: _stats(1.5, lower=2.0, upper=2.0)},
        }
        f = self._build(all_stats)
        for label in (QA1_ROW, QA2_ROW):
            self.assertEqual(self._formula(f, label), '')
            self.assertEqual(self._value(f, label), 'N/A')


class HeaderStyleTests(_LayoutBase):
    """表头行（Test Name / 参数名 …）字号为 10。"""

    def test_header_row_font_size_is_10(self):
        all_stats = {
            'FT': {PARAM: _stats(1.0)},
            'QA1': {PARAM: _stats(1.0)},
            'QA2': {PARAM: _stats(1.0)},
        }
        f = self._build(all_stats)
        ws = _load_workbook(f)[SHEET]
        for ref in ('A3', 'B3', 'C3'):
            self.assertEqual(ws[ref].font.size, 10, ref)
