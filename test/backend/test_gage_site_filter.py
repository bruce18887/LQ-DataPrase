"""Gage 工位（Site）筛选：只导出所选工位的数据，缺该工位时报错。

槽位 = 工位编号（S1→Site 1 … S8→Site 8）。同一文件可分配到多个槽位以对比
文件内的多个工位——这正是把「按文件」改成「按文件内工位」的核心场景：
一个含 Site 1~4 的文件，分配给槽位 1 与槽位 2，导出两个只含各自工位的工作表。
"""
import io
import os
from unittest import mock

import pandas as pd
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.datafiles.models import DataFile
from apps.gage.site_selection import available_sites, normalize_site_value, select_site

User = get_user_model()
GAGE_URL = '/api/v1/gage/generate_summary/'

MULTI_DF = pd.DataFrame({
    'Site_No': ['1', '1', '2', '2'],
    'T1': [1.0, 2.0, 10.0, 20.0],
})
MULTI_META = {'format': 'CTA8290D', 'mins': {'T1': '0'},
              'maxs': {'T1': '100'}, 'units': {}}


def _response_bytes(resp):
    if hasattr(resp, 'streaming_content'):
        return b''.join(resp.streaming_content)
    return resp.content


class SiteSelectionUnitTests(APITestCase):
    """site_selection 纯函数。"""

    def test_normalize_site_value(self):
        self.assertEqual(normalize_site_value('2'), '2')
        self.assertEqual(normalize_site_value(2), '2')
        self.assertEqual(normalize_site_value(2.0), '2')
        self.assertEqual(normalize_site_value(' 3 '), '3')
        self.assertEqual(normalize_site_value(float('nan')), '')

    def test_select_site_returns_only_target_rows(self):
        filtered, site_col, available = select_site(MULTI_DF, 1)
        self.assertEqual(site_col, 'Site_No')
        self.assertEqual(available, ['1', '2'])
        self.assertEqual(list(filtered['T1']), [1.0, 2.0])

    def test_select_site_missing_column(self):
        filtered, site_col, available = select_site(MULTI_DF[['T1']], 1)
        self.assertIsNone(site_col)
        self.assertIsNone(filtered)
        self.assertEqual(available, [])

    def test_available_sites_numeric_order(self):
        df = pd.DataFrame({'Site #': ['10', '2', '1']})
        self.assertEqual(available_sites(df, 'Site #'), ['1', '2', '10'])


class GageSiteFilterEndpointTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='gagesite', password='pw')
        self.client.force_authenticate(self.user)
        self.file = DataFile.objects.create(
            owner=self.user, filename='multi.csv', file_path='multi.csv',
            file_size=1, format_type='CTA8290D', status='ready',
        )

    def _post(self, assignments, **extra):
        payload = {'assignments': assignments, 'only_bin1': False,
                   'ignore_no_limit': False}
        payload.update(extra)
        with mock.patch('apps.gage.views.get_cached_parsed_file',
                        return_value=(MULTI_DF.copy(), MULTI_META, 'CTA8290D')):
            return self.client.post(GAGE_URL, payload, format='json')

    def test_same_file_two_sites_exports_filtered_sheets(self):
        # 同一文件分配给槽位 1、2 → 两个工作表，各只含对应工位数据
        resp = self._post([{'file_id': self.file.id, 'site': 1},
                           {'file_id': self.file.id, 'site': 2}])
        self.assertEqual(resp.status_code, 200, getattr(resp, 'data', ''))
        wb = load_workbook(io.BytesIO(_response_bytes(resp)))
        self.assertIn('multi_S1', wb.sheetnames)
        self.assertIn('multi_S2', wb.sheetnames)

        s1 = wb['multi_S1']
        self.assertEqual(float(s1['B14'].value), 1.0)  # Site # 列
        self.assertAlmostEqual(float(s1['J14'].value), 1.0)  # T1 第一行（I 列是 Site_No 数据列）
        self.assertAlmostEqual(float(s1['J15'].value), 2.0)
        self.assertIsNone(s1['J16'].value)  # 仅 2 行（该工位）

        s2 = wb['multi_S2']
        self.assertEqual(float(s2['B14'].value), 2.0)
        self.assertAlmostEqual(float(s2['J14'].value), 10.0)
        self.assertAlmostEqual(float(s2['J15'].value), 20.0)

    def test_missing_site_errors_with_available_sites(self):
        resp = self._post([{'file_id': self.file.id, 'site': 1},
                           {'file_id': self.file.id, 'site': 9}])
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data['error'], 'site_not_found')
        self.assertIn('Site 9', resp.data['message'])
        self.assertEqual(resp.data['details'][0]['available_sites'], ['1', '2'])

    def test_missing_site_column_errors(self):
        no_site_df = MULTI_DF[['T1']].copy()
        with mock.patch('apps.gage.views.get_cached_parsed_file',
                        return_value=(no_site_df, MULTI_META, 'CTA8290D')):
            resp = self.client.post(GAGE_URL, {
                'assignments': [{'file_id': self.file.id, 'site': 1},
                                {'file_id': self.file.id, 'site': 2}],
                'only_bin1': False, 'ignore_no_limit': False}, format='json')
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data['error'], 'no_site_column')

    def test_duplicate_assignment_rejected(self):
        resp = self._post([{'file_id': self.file.id, 'site': 1},
                           {'file_id': self.file.id, 'site': 1}])
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data['error'], 'duplicate_assignment')

    def test_fewer_than_two_assignments_rejected(self):
        resp = self._post([{'file_id': self.file.id, 'site': 1}])
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data['error'], 'need_at_least_2_files')

    def test_invalid_assignment_rejected(self):
        resp = self._post([{'file_id': self.file.id, 'site': 1},
                           {'file_id': self.file.id}])
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data['error'], 'invalid_assignment')
